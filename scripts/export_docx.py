"""Превращает markdown-отчёты в один документ Word.

Берёт reports/launch_recommendations.md и reports/summary.md и собирает
reports/Отчёт_по_рынку.docx — читаемый файл, который открывается двойным кликом
и годится, чтобы отправить его клиенту или партнёру.

Запуск:
    python scripts/export_docx.py
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import load_config, resolve, save_document  # noqa: E402

FONT = "Arial"
DARK = RGBColor(0x2F, 0x48, 0x58)
GREY = RGBColor(0x66, 0x66, 0x66)

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`(.+?)`")


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in (("Heading 1", 18, DARK), ("Heading 2", 14, DARK),
                              ("Heading 3", 12, DARK)):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)


def add_rich_text(paragraph, text: str) -> None:
    """Пишет текст, разбирая **жирный** и `моноширинный`."""
    tokens = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for token in tokens:
        if not token:
            continue
        bold = BOLD_RE.fullmatch(token)
        code = CODE_RE.fullmatch(token)
        if bold:
            run = paragraph.add_run(bold.group(1))
            run.bold = True
        elif code:
            run = paragraph.add_run(code.group(1))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            run = paragraph.add_run(token)
        run.font.name = run.font.name or FONT


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            if j >= len(table.columns):
                continue
            cell = table.cell(i, j)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            add_rich_text(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def render_markdown(doc: Document, text: str, skip_first_heading: bool = False) -> None:
    lines = text.splitlines()
    i = 0
    first_heading_done = not skip_first_heading

    while i < len(lines):
        line = lines[i].rstrip()

        # таблица
        if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if not line.strip():
            i += 1
            continue

        if line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            add_rich_text(paragraph, line[2:])
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = GREY
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            if first_heading_done:
                doc.add_heading(line[2:], level=1)
            else:
                first_heading_done = True
        elif line.startswith("---"):
            pass
        elif re.match(r"^\s*[-*] ", line):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_rich_text(paragraph, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            paragraph = doc.add_paragraph(style="List Number")
            add_rich_text(paragraph, re.sub(r"^\s*\d+\. ", "", line))
        else:
            paragraph = doc.add_paragraph()
            add_rich_text(paragraph, line)
        i += 1


def main() -> None:
    parser = argparse.ArgumentParser(description="Сборка отчёта в формате Word")
    parser.add_argument("--config", default="config.yaml")
    parser.add_argument("--out", default=None)
    args = parser.parse_args()

    cfg = load_config(args.config)
    reports_dir = resolve((cfg.get("output") or {}).get("reports_dir", "reports"))
    out_path = Path(args.out) if args.out else reports_dir / "Отчёт_по_рынку.docx"

    recommendations = reports_dir / "launch_recommendations.md"
    summary = reports_dir / "summary.md"
    if not recommendations.exists() and not summary.exists():
        sys.exit("Нет markdown-отчётов. Сначала запустите scripts/generate_report.py")

    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph("Анализ рынка полотенец на Wildberries")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in title.runs:
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = DARK

    subtitle = doc.add_paragraph("Исследование выдачи под запуск линейки эко-полотенец")
    for run in subtitle.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = GREY
    doc.add_paragraph()

    if recommendations.exists():
        render_markdown(doc, recommendations.read_text(encoding="utf-8"),
                        skip_first_heading=True)

    if summary.exists():
        doc.add_page_break()
        doc.add_heading("Приложение: сводка по собранным данным", level=1)
        render_markdown(doc, summary.read_text(encoding="utf-8"), skip_first_heading=True)

    out_path = save_document(out_path, doc.save)
    print(f"Готово: {out_path}")


if __name__ == "__main__":
    main()
